from pathlib import Path
import json, datetime, html

INPUT = Path("input data")
OUT = Path("chart.html")
MARKET_OPEN = "09:15:00"
MARKET_CLOSE = "15:30:00"
TIMEFRAMES = [("12sec", 12), ("1min", 60), ("3min", 180), ("5min", 300)]
RR = [2.0, 3.0, 4.0]


def parse_ts(ts):
    return datetime.datetime.strptime(ts, "%Y-%m-%d %H:%M:%S")


def sma(series, period):
    out = [None] * len(series)
    for i in range(len(series)):
        if i < period - 1:
            continue
        vals = [series[j] for j in range(i - period + 1, i + 1)]
        if any(v is None for v in vals):
            continue
        out[i] = sum(vals) / period
    return out


def load_ticks(day_dir, date):
    """Per-second ticks for a date, from NIFTY_YYYYMMDD_DASHBOARD.json (preferred) or .xlsx fallback."""
    compact = date.replace("-", "")
    js = day_dir / f"NIFTY_{compact}_DASHBOARD.json"
    if js.exists():
        raw = json.loads(js.read_text())
        ticks = [(parse_ts(t["timestamp"]), float(t["nifty"]), float(t.get("atm") or 0.0))
                 for t in raw.get("nifty_seconds", [])]
        flow = [(parse_ts(e["timestamp"]), e) for e in raw.get("flow", [])]
        flow.sort(key=lambda x: x[0])
        return ticks, flow, raw.get("expiry", "")
    xl = day_dir / f"NIFTY_{compact}.xlsx"
    if xl.exists():
        from openpyxl import load_workbook
        wb = load_workbook(xl, read_only=True)
        ws = wb.active
        head = [str(c.value or "").strip().lower() for c in next(ws.iter_rows(max_row=1))]
        i_ts, i_px, i_atm = head.index("timestamp"), head.index("nifty"), head.index("atm") if "atm" in head else None
        ticks = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not row[i_ts]:
                continue
            ts = row[i_ts]
            if isinstance(ts, datetime.datetime):
                dt = ts
            else:
                dt = parse_ts(str(ts).replace("T", " ")[:19])
            ticks.append((dt, float(row[i_px]), float(row[i_atm] or 0.0) if i_atm is not None else 0.0))
        wb.close()
        return ticks, [], ""
    raise FileNotFoundError(f"No JSON or XLSX input for {date}")


def build_candles(ticks, flow, date):
    ticks = sorted(ticks, key=lambda x: x[0])
    start = datetime.datetime.strptime(f"{date} {MARKET_OPEN}", "%Y-%m-%d %H:%M:%S")
    end = datetime.datetime.strptime(f"{date} {MARKET_CLOSE}", "%Y-%m-%d %H:%M:%S")
    step = datetime.timedelta(seconds=TF_SEC)
    t = start
    buckets = []
    while t < end:
        buckets.append(t)
        t += step
    rows = []
    tick_idx = 0
    flow_idx = 0
    last_flow = None
    for bucket_start in buckets:
        bucket_end = bucket_start + step
        prices = []
        while tick_idx < len(ticks) and ticks[tick_idx][0] < bucket_start:
            tick_idx += 1
        j = tick_idx
        while j < len(ticks) and ticks[j][0] < bucket_end:
            prices.append(ticks[j][1])
            j += 1
        if prices:
            o, h, l, c, n = prices[0], max(prices), min(prices), prices[-1], len(prices)
        else:
            o = h = l = c = 0.0
            n = 0
        while flow_idx < len(flow) and flow[flow_idx][0] < bucket_start:
            flow_idx += 1
        k = flow_idx
        lf = None
        while k < len(flow) and flow[k][0] < bucket_end:
            lf = flow[k][1]
            k += 1
        if lf is not None:
            last_flow = lf
        fe = last_flow or {}
        rows.append({
            "date": date,
            "time": bucket_start.strftime("%H:%M:%S"),
            "open": o, "high": h, "low": l, "close": c,
            "tick_count": n,
            "cum_ce": fe.get("cumulative_ce"),
            "cum_pe": fe.get("cumulative_pe"),
        })
    return rows


def features_for(rows):
    closes = [r["close"] if r["tick_count"] > 0 else None for r in rows]
    return sma(closes, 9), sma(closes, 44)


def pine_long(rows, f, i):
    s9, s44 = f
    if i < 1:
        return False
    cs, ce = i - 9, i - 1
    if cs < 0 or any(rows[j]["tick_count"] == 0 for j in range(cs, ce + 1)):
        return False
    sf, sn, sp = s9[i], s44[i], s44[i - 1]
    if sf is None or sn is None or sp is None:
        return False
    if not (sn > sp and sf > sn and rows[i]["close"] > sf):
        return False
    if not all(rows[j]["open"] < s9[j] and rows[j]["close"] < s9[j] for j in range(cs, ce + 1)):
        return False
    return True


def pine_short(rows, f, i):
    s9, s44 = f
    if i < 1:
        return False
    cs, ce = i - 9, i - 1
    if cs < 0 or any(rows[j]["tick_count"] == 0 for j in range(cs, ce + 1)):
        return False
    sf, sn, sp = s9[i], s44[i], s44[i - 1]
    if sf is None or sn is None or sp is None:
        return False
    if not (sn < sp and sf < sn and rows[i]["close"] < sf):
        return False
    if not all(rows[j]["open"] > s9[j] and rows[j]["close"] > s9[j] for j in range(cs, ce + 1)):
        return False
    return True


def oi_ok(rows, i, side):
    ce, pe = rows[i].get("cum_ce"), rows[i].get("cum_pe")
    if ce is None or pe is None:
        return False
    return (pe > ce) if side == "LONG" else (ce > pe)


def entry_cond_candle(rows, s9, i, direction, window=3):
    for j in range(i + 1, min(i + window + 1, len(rows))):
        r = rows[j]
        if r["tick_count"] == 0:
            continue
        sv = s9[j]
        if sv is None:
            continue
        if direction == "LONG" and r["close"] > sv:
            return j
        if direction == "SHORT" and r["close"] < sv:
            return j
    return None


def run_trades(rows, s9, s44, direction):
    trades = []
    open_until = -1
    for i in range(len(rows) - 1):
        if rows[i]["tick_count"] == 0:
            continue
        if direction == "LONG" and not pine_long(rows, (s9, s44), i):
            continue
        if direction == "SHORT" and not pine_short(rows, (s9, s44), i):
            continue
        if not oi_ok(rows, i, direction):
            continue
        ec = entry_cond_candle(rows, s9, i, direction)
        if ec is None:
            continue
        ae = None
        for j in range(ec + 1, len(rows)):
            if rows[j]["tick_count"] == 0:
                continue
            ae = j
            break
        if ae is None or ae <= open_until:
            continue
        ec_row, ae_row = rows[ec], rows[ae]
        entry_price = ae_row["close"]
        stop_price = ec_row["low"] if direction == "LONG" else ec_row["high"]
        risk = abs(entry_price - stop_price)
        if risk == 0:
            continue
        targets = {r: entry_price + r * risk if direction == "LONG" else entry_price - r * risk for r in RR}
        outcome = exit_idx = exit_price = None
        for j in range(ae + 1, len(rows)):
            r = rows[j]
            if r["tick_count"] == 0:
                continue
            hit_stop = r["low"] <= stop_price if direction == "LONG" else r["high"] >= stop_price
            for rr in RR:
                if outcome is not None:
                    continue
                hit_target = r["high"] >= targets[rr] if direction == "LONG" else r["low"] <= targets[rr]
                if hit_stop:
                    outcome, exit_idx, exit_price = "Stop Loss", j, stop_price
                    break
                if hit_target:
                    outcome, exit_idx, exit_price = "Target", j, targets[rr]
                    break
            if outcome is not None:
                break
        if outcome is None:
            last = len(rows) - 1
            while rows[last]["tick_count"] == 0 and last > ae:
                last -= 1
            outcome, exit_idx, exit_price = "End Of Data", last, rows[last]["close"]
        pnl = round((exit_price - entry_price) if direction == "LONG" else (entry_price - exit_price), 2)
        trades.append({
            "filter_time": rows[i]["time"],
            "entry_cond_time": ec_row["time"],
            "entry_time": ae_row["time"],
            "exit_time": rows[exit_idx]["time"],
            "entry_price": entry_price,
            "stop_price": stop_price,
            "exit_price": exit_price,
            "rr": max(RR),
            "outcome": outcome,
            "pnl": pnl,
        })
        open_until = exit_idx
    return trades


# ---- Render one (day, timeframe) as an SVG string with trade overlays ----

C_UP, C_DN = "#16a34a", "#dc2626"
SMA9_C, SMA44_C = "#3b82f6", "#ef4444"
BG, GRID, TXT = "#0b1220", "#1e293b", "#94a3b8"

ES = {
    "Target": ("T", "#f59e0b"),
    "Stop Loss": ("S", "#f43f5e"),        "End Of Data": ("E", "#94a3b8"),
}


def render_chart(date, tf_label, rows, s9, s44, trades):
    valid = [i for i, r in enumerate(rows) if r["tick_count"] > 0]
    if not valid:
        return ('<svg class="chartsvg" viewBox="0 0 1200 380" preserveAspectRatio="xMidYMid meet" '
                'xmlns="http://www.w3.org/2000/svg"><rect width="1200" height="380" fill="%s"/>'
                '<text x="600" y="190" fill="%s" font-size="18" text-anchor="middle">'
                'No in-market price data for %s (ticks are outside 09:15-15:30)</text></svg>' % (BG, TXT, date))
    n = len(rows)
    W = max(640, n * 2 + 80)
    H = 380
    PL, PR, PT, PB = 56, 12, 10, 22
    lows = [rows[i]["low"] for i in valid]
    highs = [rows[i]["high"] for i in valid]
    lo = min(lows)
    hi = max(highs)
    for arr in (s9, s44):
        for v in arr:
            if v is not None:
                lo = min(lo, v)
                hi = max(hi, v)
    pad = (hi - lo) * 0.06 or 1.0
    lo -= pad
    hi += pad
    span = hi - lo

    def xi(i):
        return PL + (i + 0.5) * (W - PL - PR) / n

    def yv(p):
        return PT + (hi - p) / span * (H - PT - PB)

    p = ['<svg class="chartsvg" viewBox="0 0 %d %d" preserveAspectRatio="xMidYMid meet" '
         'xmlns="http://www.w3.org/2000/svg">' % (W, H)]
    p.append('<rect x="0" y="0" width="%d" height="%d" fill="%s"/>' % (W, H, BG))
    for gy in range(5):
        yy = PT + gy * (H - PT - PB) / 4
        val = hi - gy * span / 4
        p.append('<line x1="%d" y1="%.1f" x2="%d" y2="%.1f" stroke="%s" stroke-width="1"/>'
                 % (PL, yy, W - PR, yy, GRID))
        p.append('<text x="%d" y="%.1f" fill="%s" font-size="11" text-anchor="end">%s</text>'
                 % (PL - 6, yy + 4, TXT, f"{val:.1f}"))
    for step in range(0, n, max(1, n // 8)):
        xx = xi(step)
        p.append('<line x1="%.1f" y1="%d" x2="%.1f" y2="%d" stroke="%s" stroke-width="1"/>'
                 % (xx, PT, xx, H - PB, GRID))
        p.append('<text x="%.1f" y="%d" fill="%s" font-size="11" text-anchor="middle">%s</text>'
                 % (xx, H - 7, TXT, rows[step]["time"][:5]))
    bw = max(1.0, (W - PL - PR) / n * 0.55)
    for i in valid:
        r = rows[i]
        x = xi(i)
        col = C_UP if r["close"] >= r["open"] else C_DN
        p.append('<line x1="%.1f" y1="%.1f" x2="%.1f" y2="%.1f" stroke="%s" stroke-width="1"/>'
                 % (x, yv(r["high"]), x, yv(r["low"]), col))
        y1 = yv(max(r["open"], r["close"]))
        y2 = yv(min(r["open"], r["close"]))
        p.append('<rect x="%.1f" y="%.1f" width="%.1f" height="%.1f" fill="%s"/>'
                 % (x - bw / 2, y1, bw, max(1.0, y2 - y1), col))

    def sma_path(arr, stroke):
        pts = ['%0.1f,%0.1f' % (xi(i), yv(v)) for i, v in enumerate(arr) if v is not None and rows[i]["tick_count"] > 0]
        if len(pts) < 2:
            return ""
        return '<polyline points="%s" fill="none" stroke="%s" stroke-width="1.4"/>' % (" ".join(pts), stroke)

    p.append(sma_path(s9, SMA9_C))
    p.append(sma_path(s44, SMA44_C))
    for t in trades:
        et = t["entry_time"]
        xt = t["exit_time"]
        ei = next((k for k, r in enumerate(rows) if r["time"] == et), None)
        xi2 = next((k for k, r in enumerate(rows) if r["time"] == xt), None)
        if ei is None or xi2 is None:
            continue
        x_e, x_x = xi(ei), xi(xi2)
        col = C_UP if t["direction"] == "LONG" else C_DN
        y_e = yv(t["entry_price"])
        p.append('<line x1="%.1f" y1="%.1f" x2="%.1f" y2="%.1f" stroke="%s" stroke-width="1.4" stroke-dasharray="5,3"/>'
                 % (x_e, y_e, x_x, yv(t["exit_price"]), col))
        tag, tcol = ES.get(t["outcome"], ("?", TXT))
        p.append('<circle cx="%.1f" cy="%.1f" r="4" fill="%s"/>' % (x_e, y_e, col))
        p.append('<circle cx="%.1f" cy="%.1f" r="4" fill="none" stroke="%s" stroke-width="1.4"/>'
                 % (x_x, yv(t["exit_price"]), tcol))
        p.append('<text x="%.1f" y="%.1f" fill="%s" font-size="12" font-weight="bold" text-anchor="middle">%s</text>'
                 % (x_x, yv(t["exit_price"]) - 8, tcol, tag))
    p.append('</svg>')
    return "".join(p)


def trade_row_html(t):
    cls = "up" if t["pnl"] >= 0 else "dn"
    return ("<tr><td>%s</td><td>%s</td><td class='%s'>%s</td><td>%s</td><td>%s</td><td>%s</td>"
            "<td>%.2f</td><td>%.2f</td><td>%.2f</td><td>%s</td><td class='%s'>%+.2f</td></tr>"
            % (t["date"], t["tf"], t["direction"].lower(), t["direction"], t["filter_time"],
               t["entry_cond_time"], t["entry_time"], t["entry_price"], t["stop_price"],
               t["exit_price"], t["outcome"], t["pnl"] >= 0 and "up" or "dn", t["pnl"]))


def main():
    global TF_SEC
    # ---- Build all days / timeframes ----
    days = []
    for path in sorted(INPUT.glob("NIFTY_*_DASHBOARD.json")):
        stem = path.stem.replace("NIFTY_", "").replace("_DASHBOARD", "")
        days.append(stem[:4] + "-" + stem[4:6] + "-" + stem[6:8])
    days = sorted(days)

    CHARTS = {}
    TRADES = []
    for date in days:
        ticks, flow, _ = load_ticks(INPUT, date)
        for tf_label, tf_sec in TIMEFRAMES:
            TF_SEC = tf_sec
            rows = build_candles(ticks, flow, date)
            s9, s44 = features_for(rows)
            day_trades = []
            for direction in ("LONG", "SHORT"):
                for t in run_trades(rows, s9, s44, direction):
                    t["date"] = date
                    t["tf"] = tf_label
                    t["direction"] = direction
                    day_trades.append(t)
                    TRADES.append(t)
            CHARTS[date + "|" + tf_label] = {
                "svg": render_chart(date, tf_label, rows, s9, s44, day_trades),
                "trades": day_trades,
                "candles": len([r for r in rows if r["tick_count"] > 0]),
            }
        print("done", date, flush=True)

    TRADES.sort(key=lambda t: (t["date"], ["12sec", "1min", "3min", "5min"].index(t["tf"]), t["entry_time"]))

    top = []
    top.append("<option value='ALL'>ALL DAYS</option>")
    for d in days:
        top.append("<option value='%s'>%s</option>" % (d, d))
    tf_opts = "".join("<option value='%s'>%s</option>" % (t, t) for t in ("12sec", "1min", "3min", "5min"))

    rows_js = json.dumps({k: v["trades"] for k, v in CHARTS.items()})
    charts_js = json.dumps({k: {"svg": v["svg"], "candles": v["candles"]} for k, v in CHARTS.items()})
    all_rows = "".join(trade_row_html(t) for t in TRADES)

    page = """<!DOCTYPE html>
<html><head><meta charset="utf-8"><title>NIFTY Trade Chart</title>
<style>
body{background:#0b1220;color:#e2e8f0;font:13px/1.4 Segoe UI,Arial;margin:0}
header{display:flex;gap:12px;align-items:center;padding:8px 14px;background:#0f172a;position:sticky;top:0;z-index:5}
select{background:#1e293b;color:#e2e8f0;border:1px solid #334155;border-radius:4px;padding:4px 8px}
button{background:#1e293b;color:#e2e8f0;border:1px solid #334155;border-radius:4px;padding:4px 10px;cursor:pointer}
#sum{margin-left:auto;font-size:11px;color:#94a3b8;white-space:nowrap;overflow:hidden;text-overflow:ellipsis}
#wrap{overflow:auto;height:60vh;border-bottom:1px solid #1e293b;cursor:grab}
#wrap:active{cursor:grabbing}
#stage{transform-origin:0 0;width:100%;height:100%}
.chartsvg{width:100%;height:100%;display:block}
.daylbl{padding:4px 10px;color:#60a5fa;font-size:12px;background:#0f172a}
table{border-collapse:collapse;width:100%;font-size:12px}
th{position:sticky;top:0;background:#0f172a;padding:4px 8px;text-align:left;border-bottom:1px solid #334155}
td{padding:3px 8px;border-bottom:1px solid #1e293b}
.up{color:#16a34a}.dn{color:#dc2626}
#allwrap{display:none;padding:0 0 20px}
h3{margin:10px 14px;font-size:13px;color:#94a3b8}
details{margin:8px 14px}
summary{cursor:pointer;color:#60a5fa}
</style></head><body>
<header>
<label>Day <select id="day">__DAYOPTS__</select></label>
<label>TF <select id="tf">__TFOPTS__</select></label>
<button onclick="zoom(1.25)">+</button><button onclick="zoom(0.8)">-</button>
<button onclick="fit()">fit</button>
<span id="sum"></span>
</header>
<div id="wrap"><div id="stage"></div></div>
<details id="all-details"><summary id="dsum"></summary><div id="tbl"></div></details>
<details><summary>All trades (__NALL__)</summary><div id="allwrap" style="display:block">
<table><thead><tr><th>Date</th><th>TF</th><th>Dir</th><th>Filter</th><th>Cond</th><th>Entry</th>
<th>Entry$</th><th>Stop$</th><th>Exit$</th><th>Outcome</th><th>P&amp;L</th></tr></thead>
<tbody>__ALLROWS__</tbody></table></div></details>
<script>
var CHARTS=__CHARTS__,TRADES=__TRADES__,DAYS=__DAYS__;
var scale=1,tx=0,ty=0;
function key(){return day.value+"|"+tf.value}
function draw(){
 var tfv=tf.value,ts=[],cd=0,parts=[];
 if(day.value==="ALL"){
   DAYS.forEach(function(d){
     var c=CHARTS[d+"|"+tfv];if(!c)return;
     parts.push("<div class='daylbl'>"+d+"</div>"+c.svg);
     cd+=c.candles;(TRADES[d+"|"+tfv]||[]).forEach(function(t){ts.push(t)});
   });
 } else {
   var c=CHARTS[day.value+"|"+tfv]||{svg:"",trades:[],candles:0};
   parts.push(c.svg);cd=c.candles;ts=TRADES[day.value+"|"+tfv]||[];
 }
 stage.innerHTML=parts.join("");
 var w=0;
 for(var i=0;i<ts.length;i++)w+=ts[i].pnl;
 var wins=ts.filter(function(t){return t.pnl>0}).length;
 sum.textContent="candles "+cd+" | trades "+ts.length+" | wins "+wins+" | losses "+(ts.length-wins)+" | P&L "+(w>=0?"+":"")+w.toFixed(2);
 dsum.textContent="Trades on "+(day.value==="ALL"?"ALL DAYS":day.value)+" "+tfv+" ("+ts.length+")";
 tbl.innerHTML=ts.length?"<table><thead><tr><th>Date</th><th>TF</th><th>Dir</th><th>Filter</th><th>Cond</th><th>Entry</th><th>Entry$</th><th>Stop$</th><th>Exit$</th><th>Outcome</th><th>P&L</th></tr></thead><tbody>"+
  ts.map(function(t){return "<tr><td>"+t.date+"</td><td>"+t.tf+"</td><td class='"+t.direction.toLowerCase()+"'>"+t.direction+
   "</td><td>"+t.filter_time+"</td><td>"+t.entry_cond_time+"</td><td>"+t.entry_time+"</td><td>"+t.entry_price.toFixed(2)+
   "</td><td>"+t.stop_price.toFixed(2)+"</td><td>"+t.exit_price.toFixed(2)+"</td><td>"+t.outcome+
   "</td><td class='"+(t.pnl>=0?"up":"dn")+"'>"+(t.pnl>=0?"+":"")+t.pnl.toFixed(2)+"</td></tr>"}).join("")+
  "</tbody></table>":"<p style='padding:8px 14px;color:#64748b'>No trades on this selection.</p>";
 apply();
}
function apply(){stage.style.transform="translate("+tx+"px,"+ty+"px) scale("+scale+")"}
function zoom(f){scale=Math.min(20,Math.max(0.5,scale*f));apply()}
function fit(){scale=1;tx=0;ty=0;apply()}
day.onchange=tf.onchange=draw;
wrap.onwheel=function(e){e.preventDefault();var r=wrap.getBoundingClientRect(),mx=e.clientX-r.left,my=e.clientY-r.top;
 var f=e.deltaY<0?1.15:1/1.15,ns=Math.min(20,Math.max(0.5,scale*f));tx=mx-(mx-tx)*ns/scale;ty=my-(my-ty)*ns/scale;scale=ns;apply()};
var drag=null;
wrap.onmousedown=function(e){drag={x:e.clientX,y:e.clientY,tx:tx,ty:ty}};
window.onmousemove=function(e){if(drag){tx=drag.tx+e.clientX-drag.x;ty=drag.ty+e.clientY-drag.y;apply()}};
window.onmouseup=function(){drag=null};
var dv="";
for(var i=0;i<day.options.length;i++){var cc=CHARTS[day.options[i].value+"|12sec"];if(cc&&cc.candles>0){dv=day.options[i].value;break}}
day.value=dv||day.options[1].value;draw();
document.getElementById("all-details").addEventListener("toggle",function(){wrap.style.display=this.open?"block":"none"});
</script></body></html>"""
    page = page.replace("__DAYOPTS__", "".join(top)).replace("__TFOPTS__", tf_opts)
    page = page.replace("__NALL__", str(len(TRADES))).replace("__ALLROWS__", all_rows)
    page = page.replace("__CHARTS__", charts_js).replace("__TRADES__", rows_js).replace("__DAYS__", json.dumps(days))

    OUT.write_text(page, encoding="utf-8")
    print("wrote", OUT, OUT.stat().st_size, "bytes")
    lt = sum(1 for t in TRADES if t["direction"] == "LONG")
    st = sum(1 for t in TRADES if t["direction"] == "SHORT")
    tt = sum(1 for t in TRADES if t["outcome"] == "Target")
    sl = sum(1 for t in TRADES if t["outcome"] == "Stop Loss")
    eod = sum(1 for t in TRADES if t["outcome"] == "End Of Data")
    tot = round(sum(t["pnl"] for t in TRADES), 2)
    print("trades: %d (LONG %d / SHORT %d) | Target %d / Stop %d / EOD %d | total P&L %+.2f pts"
          % (len(TRADES), lt, st, tt, sl, eod, tot))

if __name__ == "__main__":
    main()
