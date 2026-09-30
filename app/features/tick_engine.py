"""
Reference tick-to-trade engine (chart-validated).

This is the engine build_chart.py was built and verified against: per-second
ticks -> OHLC candles (09:15-15:30, zero-filled gaps flagged via
tick_count == 0), SMA9/SMA44 features, the simplified pine signal + OI
filter, the SMA9 re-cross entry condition, and the two-stage trade runner
(entry = next valid candle's close, stop = entry-condition candle low/high,
targets at RR 2/3/4, stop wins ties, End-of-data fallback, one trade at a
time).

Moved out of build_chart.py verbatim (except timeframe_seconds, which is
now an explicit parameter instead of a module global) so the chart
generator and tests share one implementation. The optimizer's analyzer
still runs its own three-stage variant (RSI-extreme + confirmation
details); migrating it here is a separate, behavior-changing decision.

The sma() below is deliberately kept as the chart's reference
implementation (naive window sum). app.features.indicators.sma is the O(n)
rolling-sum variant; on gapless input the two agree to within float
round-off (asserted in tests), but boundary comparisons like
close > sma[i] can in principle flip on that round-off, so the engine
keeps the exact arithmetic it was validated with.
"""
import datetime
import json
from pathlib import Path

from app.data.loader import MARKET_OPEN, MARKET_CLOSE

RR = [2.0, 3.0, 4.0]


def parse_ts(ts):
    return datetime.datetime.strptime(ts, "%Y-%m-%d %H:%M:%S")


def sma(series, period):
    """Reference SMA: a window containing a None (data gap) yields None."""
    out = [None] * len(series)
    for i in range(len(series)):
        if i < period - 1:
            continue
        vals = [series[j] for j in range(i - period + 1, i + 1)]
        if any(v is None for v in vals):
            continue
        out[i] = sum(vals) / period
    return out


def load_ticks(day_dir, date):
    """Per-second ticks for a date, from NIFTY_YYYYMMDD_DASHBOARD.json
    (preferred) or NIFTY_YYYYMMDD.xlsx fallback (timestamp/nifty[/atm]
    columns). Returns (ticks, flow, expiry) where flow is the 1-min OI
    series sorted by time."""
    compact = date.replace("-", "")
    js = day_dir / f"NIFTY_{compact}_DASHBOARD.json"
    if js.exists():
        raw = json.loads(js.read_text())
        ticks = [(parse_ts(t["timestamp"]), float(t["nifty"]), float(t.get("atm") or 0.0))
                 for t in raw.get("nifty_seconds", [])]
        flow = [(parse_ts(e["timestamp"]), e) for e in raw.get("flow", [])]
        flow.sort(key=lambda x: x[0])
        return ticks, flow, raw.get("expiry", "")
    xl = day_dir / f"NIFTY_{compact}.xlsx"
    if xl.exists():
        from openpyxl import load_workbook
        wb = load_workbook(xl, read_only=True)
        ws = wb.active
        head = [str(c.value or "").strip().lower() for c in next(ws.iter_rows(max_row=1))]
        i_ts, i_px, i_atm = head.index("timestamp"), head.index("nifty"), head.index("atm") if "atm" in head else None
        ticks = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not row[i_ts]:
                continue
            ts = row[i_ts]
            if isinstance(ts, datetime.datetime):
                dt = ts
            else:
                dt = parse_ts(str(ts).replace("T", " ")[:19])
            ticks.append((dt, float(row[i_px]), float(row[i_atm] or 0.0) if i_atm is not None else 0.0))
        wb.close()
        return ticks, [], ""
    raise FileNotFoundError(f"No JSON or XLSX input for {date}")


def build_candles(ticks, flow, date, timeframe_seconds):
    """Resample per-second ticks into timeframe_seconds candles covering
    09:15-15:30. Each candle joins the last flow row inside its window
    (cumulative_ce/cumulative_pe); the last seen flow row carries forward
    into later candles until a newer one appears. Gap candles are
    zero-filled with tick_count == 0."""
    ticks = sorted(ticks, key=lambda x: x[0])
    start = datetime.datetime.strptime(f"{date} {MARKET_OPEN}", "%Y-%m-%d %H:%M:%S")
    end = datetime.datetime.strptime(f"{date} {MARKET_CLOSE}", "%Y-%m-%d %H:%M:%S")
    step = datetime.timedelta(seconds=timeframe_seconds)
    t = start
    buckets = []
    while t < end:
        buckets.append(t)
        t += step
    rows = []
    tick_idx = 0
    flow_idx = 0
    last_flow = None
    for bucket_start in buckets:
        bucket_end = bucket_start + step
        prices = []
        while tick_idx < len(ticks) and ticks[tick_idx][0] < bucket_start:
            tick_idx += 1
        j = tick_idx
        while j < len(ticks) and ticks[j][0] < bucket_end:
            prices.append(ticks[j][1])
            j += 1
        if prices:
            o, h, l, c, n = prices[0], max(prices), min(prices), prices[-1], len(prices)
        else:
            o = h = l = c = 0.0
            n = 0
        while flow_idx < len(flow) and flow[flow_idx][0] < bucket_start:
            flow_idx += 1
        k = flow_idx
        lf = None
        while k < len(flow) and flow[k][0] < bucket_end:
            lf = flow[k][1]
            k += 1
        if lf is not None:
            last_flow = lf
        fe = last_flow or {}
        rows.append({
            "date": date,
            "time": bucket_start.strftime("%H:%M:%S"),
            "open": o, "high": h, "low": l, "close": c,
            "tick_count": n,
            "cum_ce": fe.get("cumulative_ce"),
            "cum_pe": fe.get("cumulative_pe"),
        })
    return rows


def features_for(rows):
    closes = [r["close"] if r["tick_count"] > 0 else None for r in rows]
    return sma(closes, 9), sma(closes, 44)


def pine_long(rows, f, i):
    s9, s44 = f
    if i < 1:
        return False
    cs, ce = i - 9, i - 1
    if cs < 0 or any(rows[j]["tick_count"] == 0 for j in range(cs, ce + 1)):
        return False
    sf, sn, sp = s9[i], s44[i], s44[i - 1]
    if sf is None or sn is None or sp is None:
        return False
    if not (sn > sp and sf > sn and rows[i]["close"] > sf):
        return False
    if not all(rows[j]["open"] < s9[j] and rows[j]["close"] < s9[j] for j in range(cs, ce + 1)):
        return False
    return True


def pine_short(rows, f, i):
    s9, s44 = f
    if i < 1:
        return False
    cs, ce = i - 9, i - 1
    if cs < 0 or any(rows[j]["tick_count"] == 0 for j in range(cs, ce + 1)):
        return False
    sf, sn, sp = s9[i], s44[i], s44[i - 1]
    if sf is None or sn is None or sp is None:
        return False
    if not (sn < sp and sf < sn and rows[i]["close"] < sf):
        return False
    if not all(rows[j]["open"] > s9[j] and rows[j]["close"] > s9[j] for j in range(cs, ce + 1)):
        return False
    return True


def oi_ok(rows, i, side):
    ce, pe = rows[i].get("cum_ce"), rows[i].get("cum_pe")
    if ce is None or pe is None:
        return False
    return (pe > ce) if side == "LONG" else (ce > pe)


def entry_cond_candle(rows, s9, i, direction, window=3):
    for j in range(i + 1, min(i + window + 1, len(rows))):
        r = rows[j]
        if r["tick_count"] == 0:
            continue
        sv = s9[j]
        if sv is None:
            continue
        if direction == "LONG" and r["close"] > sv:
            return j
        if direction == "SHORT" and r["close"] < sv:
            return j
    return None


def run_trades(rows, s9, s44, direction):
    """Two-stage trade runner over pre-built rows/features. One trade at a
    time: a setup whose entry candle falls at or before the previous
    trade's exit is skipped. Stop wins ties against same-candle targets."""
    trades = []
    open_until = -1
    for i in range(len(rows) - 1):
        if rows[i]["tick_count"] == 0:
            continue
        if direction == "LONG" and not pine_long(rows, (s9, s44), i):
            continue
        if direction == "SHORT" and not pine_short(rows, (s9, s44), i):
            continue
        if not oi_ok(rows, i, direction):
            continue
        ec = entry_cond_candle(rows, s9, i, direction)
        if ec is None:
            continue
        ae = None
        for j in range(ec + 1, len(rows)):
            if rows[j]["tick_count"] == 0:
                continue
            ae = j
            break
        if ae is None or ae <= open_until:
            continue
        ec_row, ae_row = rows[ec], rows[ae]
        entry_price = ae_row["close"]
        stop_price = ec_row["low"] if direction == "LONG" else ec_row["high"]
        risk = abs(entry_price - stop_price)
        if risk == 0:
            continue
        targets = {r: entry_price + r * risk if direction == "LONG" else entry_price - r * risk for r in RR}
        outcome = exit_idx = exit_price = None
        for j in range(ae + 1, len(rows)):
            r = rows[j]
            if r["tick_count"] == 0:
                continue
            hit_stop = r["low"] <= stop_price if direction == "LONG" else r["high"] >= stop_price
            for rr in RR:
                if outcome is not None:
                    continue
                hit_target = r["high"] >= targets[rr] if direction == "LONG" else r["low"] <= targets[rr]
                if hit_stop:
                    outcome, exit_idx, exit_price = "Stop Loss", j, stop_price
                    break
                if hit_target:
                    outcome, exit_idx, exit_price = "Target", j, targets[rr]
                    break
            if outcome is not None:
                break
        if outcome is None:
            last = len(rows) - 1
            while rows[last]["tick_count"] == 0 and last > ae:
                last -= 1
            outcome, exit_idx, exit_price = "End Of Data", last, rows[last]["close"]
        pnl = round((exit_price - entry_price) if direction == "LONG" else (entry_price - exit_price), 2)
        trades.append({
            "filter_time": rows[i]["time"],
            "entry_cond_time": ec_row["time"],
            "entry_time": ae_row["time"],
            "exit_time": rows[exit_idx]["time"],
            "entry_price": entry_price,
            "stop_price": stop_price,
            "exit_price": exit_price,
            "rr": max(RR),
            "outcome": outcome,
            "pnl": pnl,
        })
        open_until = exit_idx
    return trades
